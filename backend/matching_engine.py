"""
محرك المطابقة الفوري الشامل والمنظم (High-Performance Vectorized Matching Engine)
- ترتيب الأعمدة الدقيق وفق توجيهات المستخدم:
  1. حالة المطابقة
  2. اللوحة من ملف الداتا
  3. اللوحة من ملف الإحالة
  4. بيانات الداتا: النوع/الموديل، الملاحظات، الشارع، الحي، التاريخ
  5. بيانات الإحالة: صانع وطراز المركبة (النوع)، سنة الصنع، اسم العميل، اللون، نوع اللوحة
  6. الموقع أو الرابط من ملف الداتا
  7. باقي الأعمدة المتبقية من كلا الملفين
- المطابقة الصارمة: By Plate Only (بعد التطبيع الدقيق)
- دعم تكرار اللوحات (1-to-Many): إرجاع كافة السجلات المرتبطة باللوحة
- تصدير النتائج بالكامل إلى Excel منسق واحترافي أو CSV بترميز UTF-8 سليم
"""

import time
import json
import os
from pathlib import Path
from typing import Dict, Any, List, Optional
import duckdb
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.config import EXPORTS_DIR
from backend.normalization import normalize_plate, DIGIT_TRANSLATION, CLEANUP_PATTERN
from backend.database import get_user_duckdb_path, detect_columns


def order_matching_columns(veh_cols: List[str], ref_cols: List[str]) -> List[str]:
    """
    ترتيب أعمدة جدول النتائج وفق الترتيب المنطقي الدقيق المطلوب:
    1. حالة المطابقة
    2. اللوحة من ملف الداتا
    3. اللوحة من ملف الإحالة
    4. بيانات الداتا: النوع/الموديل، الملاحظات، الشارع، الحي، التاريخ
    5. بيانات الإحالة: صانع وطراز المركبة (النوع)، سنة الصنع، اسم العميل، اللون، نوع اللوحة
    6. الموقع أو الرابط من ملف الداتا
    7. باقي أعمدة الداتا المتبقية (إن وجدت)
    8. باقي أعمدة الإحالة المتبقية (إن وجدت)
    """
    ordered = []
    used = set()

    def add(c):
        if c and c not in used:
            ordered.append(c)
            used.add(c)

    def find_cols(col_list, keywords, exclude_keywords=None):
        matched = []
        for c in col_list:
            if c in used:
                continue
            name = c.replace('[الإحالة]', '').replace('[السيارة]', '').strip().lower()
            if exclude_keywords and any(ex in name for ex in exclude_keywords):
                continue
            if any(kw in name for kw in keywords):
                matched.append(c)
        return matched

    # 1. حالة المطابقة
    ordered.append("حالة المطابقة")
    used.add("حالة المطابقة")

    # 2. اللوحة من ملف الداتا
    for c in find_cols(veh_cols, ['لوحة', 'plate', 'اللوحه']):
        add(c)
        break

    # 3. اللوحة من ملف الإحالة
    for c in find_cols(ref_cols, ['لوحة', 'plate', 'اللوحه'], exclude_keywords=['نوع']):
        add(c)
        break

    # 4. بيانات الداتا: النوع/الموديل، الملاحظات، الشارع، الحي، التاريخ
    for c in find_cols(veh_cols, ['نوع', 'موديل', 'طراز', 'صانع', 'ماركة', 'model', 'make', 'type'], exclude_keywords=['لوحة', 'عقد']):
        add(c)
    for c in find_cols(veh_cols, ['ملاحظ', 'بيان', 'notes', 'remarks', 'وصف']):
        add(c)
    for c in find_cols(veh_cols, ['شارع', 'الشارع', 'street']):
        add(c)
    for c in find_cols(veh_cols, ['حي', 'الحي', 'district', 'neighborhood']):
        add(c)
    for c in find_cols(veh_cols, ['تاريخ', 'التاريخ', 'date', 'وقت', 'time']):
        add(c)

    # 5. بيانات الإحالة: صانع وطراز المركبة (أو النوع)، سنة الصنع، اسم العميل، اللون، نوع اللوحة
    # صانع المركبة
    for c in find_cols(ref_cols, ['صانع', 'make', 'الماركة', 'ماركة']):
        add(c)
    # طراز المركبة / الموديل / النوع
    for c in find_cols(ref_cols, ['طراز', 'نوع المركبة', 'نوع السيارة', 'الموديل', 'موديل', 'النوع', 'نوع'], exclude_keywords=['لوحة', 'لوحه', 'معامل']):
        add(c)
    # سنة الصنع / عام الصنع / الموديل
    for c in find_cols(ref_cols, ['سنة', 'سنه', 'عام الصنع', 'سنة الصنع', 'سنه الصنع', 'سنة الموديل', 'سنه الموديل', 'عام', 'year'], exclude_keywords=['معامل']):
        add(c)
    # اسم العميل
    for c in find_cols(ref_cols, ['عميل', 'اسم العميل', 'العميل', 'مالك', 'اسم المالك', 'صاحب المركبة', 'client', 'customer', 'name']):
        add(c)
    # اللون
    for c in find_cols(ref_cols, ['لون', 'اللون', 'color']):
        add(c)
    # نوع اللوحة
    for c in find_cols(ref_cols, ['نوع اللوحة', 'نوع_اللوحة', 'نوع اللوحه', 'فئة اللوحة']):
        add(c)

    # 6. الموقع أو الرابط من ملف الداتا
    for c in find_cols(veh_cols, ['موقع', 'الموقع', 'رابط', 'الرابط', 'لوكيشن', 'خريطة', 'location', 'map', 'url', 'link', 'gps', 'احداثي', 'إحداثي']):
        add(c)

    # 7. باقي أعمدة الداتا (إن وجدت، مثل الشاص، البنك، رقم العقد، إلخ)
    for c in veh_cols:
        if not c.startswith('__') and c not in used:
            add(c)

    # 8. باقي أعمدة الإحالة (إن وجدت، مثل رقم المعاملة، جهة الإحالة، إلخ)
    for c in ref_cols:
        if not c.startswith('__') and c not in used:
            add(c)

    return ordered


def match_referral_file(
    user_id: int,
    file_path: str,
    plate_col: Optional[str] = None,
    chassis_col: Optional[str] = None,
    client_col: Optional[str] = None,
    bank_col: Optional[str] = None
) -> Dict[str, Any]:
    """
    تنفيذ مطابقة ملف الإحالة ضد قاعدة بيانات المستخدم الحالية (2M سجل)
    مع الترتيب المنطقي الدقيق لكافة الأعمدة
    """
    start_time = time.time()
    db_path = get_user_duckdb_path(user_id)
    if not db_path.exists():
        raise FileNotFoundError("لم يتم العثور على قاعدة بيانات نشطة لهذا المستخدم. يرجى رفع ملف البيانات أولاً.")

    file_ext = Path(file_path).suffix.lower()

    # 1. قراءة ملف الإحالة بالكامل مع كافة الأعمدة
    if file_ext in ['.xlsx', '.xls']:
        ref_df = pd.read_excel(file_path, dtype=str, keep_default_na=False).reset_index(drop=True)
    elif file_ext in ['.csv', '.txt', '.tsv']:
        ref_df = pd.read_csv(file_path, dtype=str, keep_default_na=False, index_col=False).reset_index(drop=True)
    else:
        raise ValueError(f"صيغة ملف الإحالة غير مدعومة: {file_ext}")

    total_referral_rows = len(ref_df)
    if total_referral_rows == 0:
        return {
            "success": True,
            "total_referrals": 0,
            "matched_referrals": 0,
            "unmatched_referrals": 0,
            "total_vehicle_matches": 0,
            "multiple_matches_count": 0,
            "match_rate_percent": 0.0,
            "execution_time_ms": 0.0,
            "columns": [],
            "records": []
        }

    # تنظيف أسماء الأعمدة واكتشاف عمود اللوحة
    ref_df.columns = [str(c).strip() for c in ref_df.columns]
    detected = detect_columns(list(ref_df.columns))
    p_col = plate_col or detected['plate']

    if not p_col or p_col not in ref_df.columns:
        raise ValueError(f"تعذر تحديد عمود اللوحة في ملف الإحالة. الأعمدة المتوفرة: {list(ref_df.columns)}")

    # 2. تطبيع لوحات الإحالة وتجهيز الأعمدة مع بادئة واضحة
    norm_plates = (
        ref_df[p_col].astype(str)
        .str.normalize('NFKC')
        .str.translate(DIGIT_TRANSLATION)
        .str.replace(CLEANUP_PATTERN, '', regex=True)
        .str.lower()
        .str.strip()
    )

    clean_ref_df = ref_df.copy()
    clean_ref_df.columns = [f"[الإحالة] {c}" for c in clean_ref_df.columns]
    clean_ref_df['__ref_row_id__'] = list(range(1, total_referral_rows + 1))
    clean_ref_df['__ref_plate_norm__'] = norm_plates

    # 3. الربط الفوري عبر DuckDB وترتيب الأعمدة بدقة
    con = duckdb.connect(str(db_path))
    try:
        con.register("ref_table", clean_ref_df)

        con.execute("DROP TABLE IF EXISTS temp_match;")
        con.execute("""
            CREATE TABLE temp_match AS 
            SELECT 
                r.__ref_row_id__ AS "__row_id__",
                CASE 
                    WHEN v.__veh_plate_norm__ IS NULL THEN 'غير متطابق'
                    WHEN COUNT(v.__veh_plate_norm__) OVER (PARTITION BY r.__ref_row_id__) > 1 THEN 'تطابق متعدد'
                    ELSE 'متطابق'
                END AS "حالة المطابقة",
                r.* EXCLUDE (__ref_row_id__, __ref_plate_norm__),
                v.* EXCLUDE (__veh_id__, __veh_plate_norm__),
                CASE WHEN v.__veh_plate_norm__ IS NOT NULL THEN 1 ELSE 0 END AS "__is_matched__"
            FROM ref_table r
            LEFT JOIN vehicles v ON r.__ref_plate_norm__ = v.__veh_plate_norm__ AND r.__ref_plate_norm__ != ''
            ORDER BY r.__ref_row_id__;
        """)
        con.unregister("ref_table")

        # ترتيب الأعمدة حسب المواصفات المطلوبة
        all_raw_cols = [c[0] for c in con.execute('DESCRIBE temp_match').fetchall()]
        veh_cols = [c for c in all_raw_cols if c.startswith('[السيارة]')]
        ref_cols = [c for c in all_raw_cols if c.startswith('[الإحالة]')]
        ordered_cols = order_matching_columns(veh_cols, ref_cols)

        # إنشاء جدول last_match بالأعمدة المرتبة بدقة
        cols_sql = ', '.join([f'"{c}"' for c in ordered_cols])
        con.execute("DROP TABLE IF EXISTS last_match;")
        con.execute(f'CREATE TABLE last_match AS SELECT "__row_id__", "__is_matched__", {cols_sql} FROM temp_match;')
        con.execute("DROP TABLE IF EXISTS temp_match;")

        # 4. حساب الإحصائيات الدقيقة
        stats_row = con.execute("""
            SELECT 
                COUNT(*) AS total_output_rows,
                COALESCE(SUM(__is_matched__), 0) AS total_vehicle_matches,
                COUNT(DISTINCT CASE WHEN __is_matched__ = 1 THEN __row_id__ END) AS matched_referrals_count
            FROM last_match;
        """).fetchone()

        total_output_rows = stats_row[0]
        total_vehicle_matches = int(stats_row[1])
        matched_referrals_count = int(stats_row[2])
        unmatched_referrals_count = total_referral_rows - matched_referrals_count

        multiple_matches_count = con.execute("""
            SELECT COUNT(*) FROM (
                SELECT __row_id__ FROM last_match WHERE __is_matched__ = 1 GROUP BY __row_id__ HAVING COUNT(*) > 1
            );
        """).fetchone()[0]

    finally:
        con.close()

    execution_time_ms = round((time.time() - start_time) * 1000, 2)
    match_rate = round((matched_referrals_count / total_referral_rows * 100), 1) if total_referral_rows > 0 else 0.0

    return {
        "success": True,
        "total_referrals": total_referral_rows,
        "matched_referrals": matched_referrals_count,
        "unmatched_referrals": unmatched_referrals_count,
        "total_vehicle_matches": total_vehicle_matches,
        "multiple_matches_count": multiple_matches_count,
        "match_rate_percent": match_rate,
        "execution_time_ms": execution_time_ms,
        "total_output_rows": total_output_rows
    }


def get_match_results_page(
    user_id: int,
    page: int = 1,
    page_size: int = 50,
    filter_status: str = 'all',
    search_query: str = ''
) -> Dict[str, Any]:
    """
    استرجاع صفحة من نتائج المطابقة الأخيرة مع الأعمدة المرتبة بدقة
    """
    db_path = get_user_duckdb_path(user_id)
    if not db_path.exists():
        return {"total": 0, "pages": 0, "page": page, "columns": [], "records": []}

    con = duckdb.connect(str(db_path), read_only=True)
    try:
        table_check = con.execute("SELECT count(*) FROM information_schema.tables WHERE table_name = 'last_match';").fetchone()[0]
        if table_check == 0:
            return {"total": 0, "pages": 0, "page": page, "columns": [], "records": []}

        # جلب قائمة كافة الأعمدة المتاحة بالترتيب المحفوظ
        all_cols = [c[0] for c in con.execute('DESCRIBE last_match').fetchall() if not c[0].startswith('__')]

        where_clauses = []
        params = []

        if filter_status == 'matched':
            where_clauses.append("__is_matched__ = 1")
        elif filter_status == 'unmatched':
            where_clauses.append("__is_matched__ = 0")
        elif filter_status == 'multiple':
            where_clauses.append('"حالة المطابقة" = \'تطابق متعدد\'')

        if search_query:
            sq = search_query.strip()
            conditions = [f'CAST("{col}" AS VARCHAR) LIKE ?' for col in all_cols]
            where_clauses.append(f"({' OR '.join(conditions)})")
            params.extend([f"%{sq}%"] * len(conditions))

        where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

        count_sql = f"SELECT COUNT(*) FROM last_match {where_sql};"
        total_records = con.execute(count_sql, params).fetchone()[0]

        total_pages = (total_records + page_size - 1) // page_size if page_size > 0 else 1
        page = max(1, min(page, total_pages)) if total_pages > 0 else 1
        offset = (page - 1) * page_size

        cols_select = ', '.join([f'"{c}"' for c in all_cols])
        select_sql = f"""
            SELECT {cols_select} 
            FROM last_match 
            {where_sql}
            ORDER BY __row_id__
            LIMIT ? OFFSET ?;
        """
        page_params = params + [page_size, offset]
        df_page = con.execute(select_sql, page_params).df().fillna('')
        records = df_page.to_dict(orient='records')

        return {
            "total": total_records,
            "pages": total_pages,
            "page": page,
            "columns": all_cols,
            "records": records
        }
    finally:
        con.close()


def export_match_results(user_id: int, file_format: str = 'xlsx', filter_status: str = 'all') -> str:
    """
    تصدير نتائج المطابقة كاملة بالأعمدة المرتبة بدقة إلى Excel (.xlsx) احترافي أو CSV بترميز UTF-8
    """
    db_path = get_user_duckdb_path(user_id)
    if not db_path.exists():
        raise FileNotFoundError("لا توجد نتائج مطابقة لتصديرها.")

    con = duckdb.connect(str(db_path), read_only=True)
    try:
        table_check = con.execute("SELECT count(*) FROM information_schema.tables WHERE table_name = 'last_match';").fetchone()[0]
        if table_check == 0:
            raise FileNotFoundError("لا توجد نتائج مطابقة مسجلة.")

        all_cols = [c[0] for c in con.execute('DESCRIBE last_match').fetchall() if not c[0].startswith('__')]
        cols_select = ', '.join([f'"{c}"' for c in all_cols])

        where_sql = ""
        if filter_status == 'matched':
            where_sql = "WHERE __is_matched__ = 1"
        elif filter_status == 'unmatched':
            where_sql = "WHERE __is_matched__ = 0"
        elif filter_status == 'multiple':
            where_sql = 'WHERE "حالة المطابقة" = \'تطابق متعدد\''

        query = f"SELECT {cols_select} FROM last_match {where_sql} ORDER BY __row_id__;"
        df = con.execute(query).df().fillna('')
    finally:
        con.close()

    timestamp = time.strftime("%Y%m%d_%H%M%S")

    if file_format == 'csv':
        out_filename = f"matching_result_ordered_{user_id}_{timestamp}.csv"
        out_path = EXPORTS_DIR / out_filename
        df.to_csv(out_path, index=False, encoding='utf-8-sig')
        return str(out_path)

    # تصدير Excel منسق واحترافي
    out_filename = f"matching_result_ordered_{user_id}_{timestamp}.xlsx"
    out_path = EXPORTS_DIR / out_filename

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "نتائج المطابقة المرتبة"
    ws.views.sheetView[0].rightToLeft = True

    headers = list(df.columns)
    ws.append(headers)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    matched_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    unmatched_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    multiple_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10)

    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        status = str(row[0])
        row_fill = matched_fill if status == 'متطابق' else (multiple_fill if status == 'تطابق متعدد' else unmatched_fill)

        for col_idx, value in enumerate(row, 1):
            val_str = "" if pd.isna(value) else str(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=val_str)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx == 1:
                cell.fill = row_fill
                cell.font = Font(name="Segoe UI", size=10, bold=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    ws.row_dimensions[1].height = 28
    wb.save(out_path)
    return str(out_path)


def quick_search_single_plate(user_id: int, query_plate: str) -> List[Dict[str, Any]]:
    """
    البحث الفوري عن لوحة واحدة وإرجاع كافة أعمدة السيارة بالكامل
    """
    db_path = get_user_duckdb_path(user_id)
    if not db_path.exists():
        return []

    norm_query = normalize_plate(query_plate)
    if not norm_query:
        return []

    con = duckdb.connect(str(db_path), read_only=True)
    try:
        table_check = con.execute("SELECT count(*) FROM information_schema.tables WHERE table_name = 'vehicles';").fetchone()[0]
        if table_check == 0:
            return []

        query = "SELECT * EXCLUDE (__veh_id__, __veh_plate_norm__) FROM vehicles WHERE __veh_plate_norm__ = ?;"
        df = con.execute(query, [norm_query]).df().fillna('')
        return df.to_dict(orient='records')
    finally:
        con.close()
